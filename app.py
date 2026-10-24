import hashlib
import html
import io
import json
import re
import zipfile
from pathlib import Path

import requests
import streamlit as st
from docx import Document
from pypdf import PdfReader


# =========================================================
# CONFIGURATION
# =========================================================

st.set_page_config(
    page_title="LazyDesk AI — Smart Document Workspace",
    page_icon="🗂️",
    layout="wide",
    initial_sidebar_state="expanded",
)

OLLAMA_URL = "http://localhost:11434"
MODEL = "gemma3:1b"

MAX_FILE_SIZE = 10 * 1024 * 1024
MAX_PDF_PAGES = 100
MAX_TEXT_LENGTH = 10_000

CATEGORIES = [
    "Computer Science",
    "Education",
    "Business",
    "Finance",
    "Research",
    "Personal",
    "Other",
]


# =========================================================
# DESIGN SYSTEM
# =========================================================

st.markdown(
    """
    <style>
    @import url('https://fonts.googleapis.com/css2?family=DM+Sans:wght@400;500;600;700&family=Manrope:wght@400;500;600;700;800&display=swap');

    :root {
        --ink: #182039;
        --muted: #788198;
        --purple: #6655e8;
        --purple-dark: #5141d0;
        --line: #e7e9f2;
        --surface: #ffffff;
        --canvas: #f7f8fc;
    }

    html, body, [class*="css"] {
        font-family: 'DM Sans', sans-serif;
    }

    .stApp {
        background: var(--canvas);
        color: var(--ink);
    }

    header[data-testid="stHeader"] {
        background: rgba(247, 248, 252, 0.9);
    }

    [data-testid="stSidebar"] {
        background: #ffffff;
        border-right: 1px solid var(--line);
    }

    [data-testid="stSidebar"] > div:first-child {
        padding-top: 1.3rem;
    }

    .brand {
        display: flex;
        align-items: center;
        gap: 12px;
        padding: 8px 0 18px;
    }

    .brand-icon {
        display: flex;
        align-items: center;
        justify-content: center;
        width: 46px;
        height: 46px;
        border-radius: 15px;
        background: linear-gradient(135deg, #7968ff, #5745d6);
        color: white;
        font-size: 23px;
        box-shadow: 0 8px 20px rgba(102, 85, 232, .20);
    }

    .brand-name {
        font-family: 'Manrope', sans-serif;
        color: var(--ink);
        font-size: 21px;
        font-weight: 800;
        letter-spacing: -0.7px;
    }

    .brand-sub {
        color: var(--muted);
        font-size: 12px;
        margin-top: 1px;
    }

    .hero {
        position: relative;
        overflow: hidden;
        background:
            radial-gradient(ellipse at 95% 0%, rgba(188, 176, 255, .35), transparent 40%),
            linear-gradient(125deg, #ffffff 0%, #f5f2ff 100%);
        border: 1px solid #e7e0ff;
        border-radius: 25px;
        padding: 35px 36px 31px;
        margin: 5px 0 24px;
        box-shadow: 0 10px 35px rgba(44, 37, 102, .035);
    }

    .eyebrow {
        display: inline-flex;
        align-items: center;
        gap: 7px;
        background: #ffffff;
        border: 1px solid #e4ddff;
        border-radius: 100px;
        padding: 7px 12px;
        color: #6551dc;
        font-size: 11px;
        font-weight: 800;
        letter-spacing: 1px;
        text-transform: uppercase;
    }

    .hero h1 {
        font-family: 'Manrope', sans-serif;
        font-size: clamp(30px, 4vw, 46px);
        line-height: 1.15;
        letter-spacing: -1.8px;
        color: #171d39;
        font-weight: 800;
        margin: 23px 0 12px;
    }

    .hero h1 span {
        color: #6955e8;
    }

    .hero-copy {
        max-width: 640px;
        font-size: 15px;
        line-height: 1.8;
        color: #717b96;
        margin: 0;
    }

    .hero-bottom {
        display: flex;
        gap: 9px;
        flex-wrap: wrap;
        margin-top: 24px;
    }

    .hero-chip {
        border: 1px solid #e5e5f2;
        background: rgba(255,255,255,.8);
        border-radius: 9px;
        padding: 8px 11px;
        font-size: 12px;
        font-weight: 600;
        color: #505a75;
    }

    .section-kicker {
        color: #6a58e6;
        font-size: 11px;
        font-weight: 800;
        letter-spacing: 1.5px;
        text-transform: uppercase;
        margin-bottom: 6px;
    }

    .section-title {
        font-family: 'Manrope', sans-serif;
        font-size: 25px;
        font-weight: 800;
        letter-spacing: -.8px;
        color: #1c2340;
        margin: 0 0 6px;
    }

    .section-copy {
        color: #7b849a;
        font-size: 14px;
        line-height: 1.7;
        margin-bottom: 18px;
    }

    .feature-card {
        background: #ffffff;
        border: 1px solid var(--line);
        border-radius: 17px;
        padding: 19px 18px;
        height: 100%;
        min-height: 145px;
        box-shadow: 0 4px 14px rgba(30, 38, 75, .025);
    }

    .feature-icon {
        display: flex;
        align-items: center;
        justify-content: center;
        width: 39px;
        height: 39px;
        border-radius: 12px;
        background: #f0edff;
        font-size: 19px;
        margin-bottom: 13px;
    }

    .feature-title {
        font-size: 14px;
        font-weight: 800;
        color: #222944;
        margin-bottom: 5px;
    }

    .feature-copy {
        color: #8089a0;
        font-size: 12px;
        line-height: 1.65;
    }

    .panel {
        background: #ffffff;
        border: 1px solid var(--line);
        border-radius: 19px;
        padding: 22px;
        box-shadow: 0 5px 18px rgba(30, 38, 75, .025);
        margin-bottom: 16px;
    }

    .panel-heading {
        font-family: 'Manrope', sans-serif;
        font-size: 18px;
        font-weight: 800;
        color: #202742;
        margin-bottom: 4px;
    }

    .panel-description {
        color: #838ba0;
        font-size: 13px;
        line-height: 1.65;
        margin-bottom: 17px;
    }

    .metric-card {
        background: #ffffff;
        border: 1px solid var(--line);
        border-radius: 15px;
        padding: 17px;
        min-height: 108px;
    }

    .metric-label {
        font-size: 12px;
        color: #81899e;
        font-weight: 600;
    }

    .metric-value {
        color: #202742;
        font-family: 'Manrope', sans-serif;
        font-size: 27px;
        font-weight: 800;
        letter-spacing: -.8px;
        margin-top: 6px;
    }

    .metric-note {
        font-size: 11px;
        color: #6b59df;
        margin-top: 2px;
    }

    .status-pill {
        display: inline-flex;
        align-items: center;
        gap: 7px;
        padding: 7px 10px;
        border-radius: 9px;
        background: #e9f9ef;
        color: #16824b;
        font-size: 12px;
        font-weight: 700;
    }

    .status-dot {
        height: 7px;
        width: 7px;
        border-radius: 50%;
        background: #1aa45e;
    }

    .file-row {
        display: flex;
        align-items: center;
        gap: 12px;
        background: #fbfbfe;
        border: 1px solid #ebedf5;
        border-radius: 12px;
        padding: 12px;
        margin: 8px 0;
    }

    .file-icon {
        display: flex;
        align-items: center;
        justify-content: center;
        width: 39px;
        height: 43px;
        background: #f0edff;
        border-radius: 10px;
        font-size: 19px;
        flex-shrink: 0;
    }

    .file-name {
        color: #252c47;
        font-size: 13px;
        font-weight: 700;
        overflow-wrap: anywhere;
    }

    .file-meta {
        color: #858da2;
        font-size: 11px;
        margin-top: 4px;
    }

    .result-box {
        background: #ffffff;
        border: 1px solid var(--line);
        border-radius: 17px;
        padding: 19px;
        margin: 12px 0;
        box-shadow: 0 5px 18px rgba(30, 38, 75, .025);
    }

    .category-pill {
        display: inline-block;
        background: #f0edff;
        color: #5e4bd2;
        padding: 6px 10px;
        border-radius: 8px;
        font-size: 11px;
        font-weight: 800;
        margin-bottom: 10px;
    }

    .result-summary {
        font-size: 13px;
        line-height: 1.8;
        color: #626c85;
        overflow-wrap: anywhere;
    }

    .result-filename {
        background: #f7f8fc;
        border: 1px solid #e9ebf4;
        padding: 10px 12px;
        border-radius: 9px;
        color: #3d4663;
        font-size: 12px;
        overflow-wrap: anywhere;
    }

    div.stButton > button,
    div.stDownloadButton > button {
        border-radius: 11px;
        font-weight: 700;
        min-height: 43px;
        transition: all .18s ease;
    }

    div.stButton > button[kind="primary"] {
        background: linear-gradient(105deg, #7662f2, #5d49dc);
        border: none;
        box-shadow: 0 5px 15px rgba(102, 85, 232, .16);
    }

    div.stButton > button[kind="primary"]:hover {
        background: linear-gradient(105deg, #6854e7, #4e3bc8);
        box-shadow: 0 8px 20px rgba(102, 85, 232, .24);
        transform: translateY(-1px);
    }

    div.stDownloadButton > button {
        border: 1px solid #ded8ff;
        color: #5947d1;
        background: #f7f5ff;
    }

    div[data-testid="stFileUploader"] {
        background: #fbfaff;
        border: 1px dashed #cfc6ff;
        border-radius: 14px;
        padding: 13px;
    }

    div[data-testid="stFileUploader"] section {
        background: transparent;
    }

    div[data-testid="stFileUploader"] button {
        border-radius: 9px;
    }

    div[data-testid="stAlert"] {
        border-radius: 12px;
    }

    div[data-testid="stExpander"] {
        background: #ffffff;
        border: 1px solid var(--line);
        border-radius: 12px;
    }

    .footer {
        border-top: 1px solid #e5e7f0;
        padding: 24px 0 8px;
        margin-top: 30px;
        text-align: center;
        color: #8991a5;
        font-size: 12px;
        line-height: 1.9;
    }

    .footer strong {
        color: #5e50c9;
    }

    @media (max-width: 768px) {
        .hero {
            padding: 25px 21px;
            border-radius: 19px;
        }

        .hero h1 {
            font-size: 32px;
        }

        .panel {
            padding: 16px;
        }
    }
    </style>
    """,
    unsafe_allow_html=True,
)


# =========================================================
# UTILITIES
# =========================================================

def safe_filename(name):
    """Remove characters that are unsafe in Windows filenames."""
    name = Path(name).name
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", name)
    name = name.strip(" .")
    return name or "document"


def file_digest(data):
    return hashlib.sha256(data).hexdigest()


def readable_size(size):
    if size < 1024:
        return f"{size} B"
    if size < 1024 * 1024:
        return f"{size / 1024:.1f} KB"
    return f"{size / (1024 * 1024):.2f} MB"


def get_file_text(name, data):
    """Extract readable text from TXT, PDF, or DOCX."""
    extension = Path(name).suffix.lower()

    if extension == ".txt":
        text = data.decode("utf-8-sig", errors="replace")

    elif extension == ".pdf":
        reader = PdfReader(io.BytesIO(data))

        if len(reader.pages) > MAX_PDF_PAGES:
            raise ValueError(
                f"PDF has {len(reader.pages)} pages. "
                f"The maximum supported is {MAX_PDF_PAGES} pages."
            )

        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as exc:
                raise ValueError(
                    "This PDF is password-protected and cannot be read."
                ) from exc

        pages = []
        for page in reader.pages:
            pages.append(page.extract_text() or "")

        text = "\n".join(pages)

        if not text.strip():
            raise ValueError(
                "No selectable text was found. This may be a scanned PDF "
                "that needs OCR."
            )

    elif extension == ".docx":
        document = Document(io.BytesIO(data))
        parts = [paragraph.text for paragraph in document.paragraphs]

        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))

        text = "\n".join(parts)

    else:
        raise ValueError("Unsupported format. Upload PDF, DOCX, or TXT.")

    text = text.strip()

    if not text:
        raise ValueError("This file does not contain readable text.")

    if len(text) > MAX_TEXT_LENGTH:
        text = text[:MAX_TEXT_LENGTH]

    return text


def check_ollama():
    """Check whether Ollama and the configured model are available."""
    try:
        response = requests.get(
            f"{OLLAMA_URL}/api/tags",
            timeout=3,
        )
        response.raise_for_status()

        models = response.json().get("models", [])
        model_names = [
            item.get("name", "")
            for item in models
        ]

        available = any(
            name == MODEL or name.startswith(MODEL + ":")
            for name in model_names
        )

        return True, available, None

    except requests.RequestException as exc:
        return False, False, str(exc)


def analyze_document(filename, text):
    """Ask the local Ollama model to organize a document."""
    prompt = f"""
You are a careful document organization assistant.

Analyze the document below and return ONLY valid JSON using this schema:
{{
  "category": "one category from the list",
  "summary": "a concise, useful summary in 2 to 4 sentences",
  "suggested_filename": "a short descriptive filename without extension"
}}

Allowed categories:
{", ".join(CATEGORIES)}

Rules:
- Use only information supported by the document.
- Never invent facts that are not in the document.
- Choose the most appropriate category.
- Make the filename descriptive and filesystem-safe.
- Do not include Markdown or text outside the JSON object.

Original filename: {filename}

Document content:
{text}
"""

    response = requests.post(
        f"{OLLAMA_URL}/api/generate",
        json={
            "model": MODEL,
            "prompt": prompt,
            "stream": False,
            "format": "json",
            "options": {
                "temperature": 0.2,
            },
        },
        timeout=(10, 300),
    )

    response.raise_for_status()
    answer = response.json().get("response", "").strip()

    try:
        result = json.loads(answer)
    except json.JSONDecodeError:
        match = re.search(r"\{.*\}", answer, re.DOTALL)
        if not match:
            raise ValueError("The AI returned an invalid response.")
        result = json.loads(match.group(0))

    category = str(result.get("category", "Other")).strip()

    if category not in CATEGORIES:
        category = "Other"

    summary = str(result.get("summary", "")).strip()
    suggested = str(result.get("suggested_filename", "")).strip()

    if not summary:
        summary = "The AI did not return a summary for this document."

    if not suggested:
        suggested = Path(filename).stem

    suggested = safe_filename(suggested)
    suggested = Path(suggested).stem

    extension = Path(filename).suffix.lower()

    return {
        "original_name": filename,
        "category": category,
        "summary": summary,
        "suggested_filename": suggested + extension,
    }


def build_zip(results, file_data):
    """Create category folders containing organized copies."""
    output = io.BytesIO()
    used_paths = set()

    with zipfile.ZipFile(
        output,
        mode="w",
        compression=zipfile.ZIP_DEFLATED,
    ) as archive:

        for digest, result in results.items():
            if digest not in file_data:
                continue

            category = safe_filename(result["category"])
            filename = safe_filename(result["suggested_filename"])

            path = f"{category}/{filename}"
            base = Path(filename).stem
            suffix = Path(filename).suffix
            counter = 2

            while path.lower() in used_paths:
                path = f"{category}/{base}_{counter}{suffix}"
                counter += 1

            used_paths.add(path.lower())
            archive.writestr(path, file_data[digest]["data"])

        manifest = {
            "application": "LazyDesk AI",
            "model": MODEL,
            "documents": [
                {
                    "original_name": result["original_name"],
                    "category": result["category"],
                    "suggested_filename": result["suggested_filename"],
                    "summary": result["summary"],
                }
                for result in results.values()
            ],
        }

        archive.writestr(
            "organization_report.json",
            json.dumps(manifest, indent=2, ensure_ascii=False),
        )

    output.seek(0)
    return output.getvalue()


# =========================================================
# SESSION STATE
# =========================================================

if "results" not in st.session_state:
    st.session_state.results = {}

if "file_data" not in st.session_state:
    st.session_state.file_data = {}

if "last_upload_signature" not in st.session_state:
    st.session_state.last_upload_signature = ""

if "workspace_notice" not in st.session_state:
    st.session_state.workspace_notice = ""


# =========================================================
# SIDEBAR
# =========================================================

with st.sidebar:
    st.markdown(
        """
        <div class="brand">
            <div class="brand-icon">✦</div>
            <div>
                <div class="brand-name">LazyDesk AI</div>
                <div class="brand-sub">Your intelligent document workspace</div>
            </div>
        </div>
        """,
        unsafe_allow_html=True,
    )

    st.divider()

    online, model_available, connection_error = check_ollama()

    if online and model_available:
        st.markdown(
            """
            <div class="status-pill">
                <span class="status-dot"></span>
                Local AI is online
            </div>
            """,
            unsafe_allow_html=True,
        )
    elif online:
        st.warning("Ollama is running, but the model was not found.")
        st.caption(f"Expected model: {MODEL}")
    else:
        st.error("Local AI is offline")
        st.caption("Start Ollama to enable document analysis.")

    st.markdown("#### ACTIVE MODEL")

    st.code(MODEL, language=None)

    if model_available:
        st.caption("Gemma 3 is ready")
    elif online:
        st.caption(f"Run `ollama pull {MODEL}` in your terminal.")

    st.divider()

    st.markdown("### 🔐 Privacy by design")
    st.caption(
        "Document analysis uses your local Ollama service. "
        "This app does not send document content to a remote AI API."
    )

    st.divider()

    st.markdown("### Supported formats")
    st.markdown(
        "- 📕 PDF documents\n"
        "- 📝 Word documents (.docx)\n"
        "- 📄 Text files (.txt)"
    )

    st.caption("Maximum file size: 10 MB per document")

    st.divider()

    st.caption("LAZYDESK AI · LOCAL-FIRST AI")


# =========================================================
# HERO
# =========================================================

st.markdown(
    """
    <section class="hero">
        <div class="eyebrow">✦ PRIVATE AI WORKSPACE</div>
        <h1>Your documents,<br><span>under control.</span></h1>
        <p class="hero-copy">
            Turn scattered files into an organized workspace.
            Let local AI categorize documents, summarize key information,
            and prepare sorted copies — all in one place.
        </p>
        <div class="hero-bottom">
            <span class="hero-chip">🔒 Private by design</span>
            <span class="hero-chip">✦ Powered by Gemma 3</span>
            <span class="hero-chip">⚡ One-click organization</span>
        </div>
    </section>
    """,
    unsafe_allow_html=True,
)


# =========================================================
# WORKSPACE METRICS
# =========================================================

st.markdown(
    """
    <div class="section-kicker">YOUR WORKSPACE</div>
    <div class="section-title">Your document command center</div>
    <div class="section-copy">
        Upload your documents, let your local AI do the organizing,
        and export a clean, structured collection.
    </div>
    """,
    unsafe_allow_html=True,
)

total_files = len(st.session_state.file_data)
analyzed_files = len(st.session_state.results)
category_count = len({
    result["category"]
    for result in st.session_state.results.values()
})

metric_columns = st.columns(4)

metrics = [
    ("Documents loaded", total_files, "Files in this session"),
    ("AI analyzed", analyzed_files, "Completed analyses"),
    ("Categories found", category_count, "Distinct categories"),
    ("AI model", "Gemma 3", "Local inference"),
]

for column, metric in zip(metric_columns, metrics):
    label, value, note = metric

    with column:
        st.markdown(
            f"""
            <div class="metric-card">
                <div class="metric-label">{html.escape(str(label))}</div>
                <div class="metric-value">{html.escape(str(value))}</div>
                <div class="metric-note">{html.escape(str(note))}</div>
            </div>
            """,
            unsafe_allow_html=True,
        )


st.markdown("")


# =========================================================
# FEATURE CARDS
# =========================================================

feature_columns = st.columns(3)

features = [
    (
        "🧠",
        "Smart categorization",
        "Automatically group documents into meaningful categories.",
    ),
    (
        "✦",
        "AI-powered summaries",
        "Understand important points without reading everything.",
    ),
    (
        "📦",
        "One-click organization",
        "Download categorized copies and a JSON report in one ZIP.",
    ),
]

for column, (emoji, title, description) in zip(feature_columns, features):
    with column:
        st.markdown(
            f"""
            <div class="feature-card">
                <div class="feature-icon">{emoji}</div>
                <div class="feature-title">{html.escape(title)}</div>
                <div class="feature-copy">{html.escape(description)}</div>
            </div>
            """,
            unsafe_allow_html=True,
        )


st.markdown("")


# =========================================================
# UPLOAD WORKSPACE
# =========================================================

st.markdown(
    """
    <div class="section-kicker">WORKSPACE</div>
    <div class="section-title">Bring your files together</div>
    <div class="section-copy">
        Upload PDF, Word, or text documents. Review the selected files,
        analyze them with your local model, then export organized copies.
    </div>
    """,
    unsafe_allow_html=True,
)

with st.container(border=True):
    st.markdown(
        '<div class="panel-heading">📂 Upload documents</div>',
        unsafe_allow_html=True,
    )

    st.markdown(
        '<div class="panel-description">'
        'Choose multiple documents at once. Your original files will not be modified.'
        '</div>',
        unsafe_allow_html=True,
    )

    uploaded_files = st.file_uploader(
        "Drop documents here or browse your files",
        type=["pdf", "docx", "txt"],
        accept_multiple_files=True,
        help="Supported formats: PDF, DOCX and TXT. Maximum 10 MB per file.",
        key="lazy_upload",
    )

    upload_signature_parts = []

    for uploaded_file in uploaded_files or []:
        data = uploaded_file.getvalue()
        upload_signature_parts.append(
            f"{uploaded_file.name}:{file_digest(data)}"
        )

    upload_signature = "|".join(sorted(upload_signature_parts))

    if upload_signature != st.session_state.last_upload_signature:
        st.session_state.last_upload_signature = upload_signature
        st.session_state.workspace_notice = ""

        current_data = {}
        current_results = {}

        for uploaded_file in uploaded_files or []:
            data = uploaded_file.getvalue()
            digest = file_digest(data)

            current_data[digest] = {
                "name": safe_filename(uploaded_file.name),
                "data": data,
                "size": len(data),
            }

            if digest in st.session_state.results:
                current_results[digest] = st.session_state.results[digest]

        st.session_state.file_data = current_data
        st.session_state.results = current_results

    if uploaded_files:
        st.markdown("#### Selected files")

        total_size = sum(
            len(item.getvalue())
            for item in uploaded_files
        )

        info_columns = st.columns(3)

        info_columns[0].metric("Files selected", len(uploaded_files))
        info_columns[1].metric("Combined size", readable_size(total_size))
        info_columns[2].metric(
            "Ready for AI",
            sum(
                len(item.getvalue()) <= MAX_FILE_SIZE
                for item in uploaded_files
            ),
        )

        for uploaded_file in uploaded_files:
            extension = Path(uploaded_file.name).suffix.lower()
            emoji = {
                ".pdf": "📕",
                ".docx": "📝",
                ".txt": "📄",
            }.get(extension, "📁")

            st.markdown(
                f"""
                <div class="file-row">
                    <div class="file-icon">{emoji}</div>
                    <div style="flex:1; min-width:0;">
                        <div class="file-name">
                            {html.escape(uploaded_file.name)}
                        </div>
                        <div class="file-meta">
                            {html.escape(extension.upper().replace(".", ""))}
                            document · {readable_size(len(uploaded_file.getvalue()))}
                        </div>
                    </div>
                </div>
                """,
                unsafe_allow_html=True,
            )

        oversized = [
            item.name
            for item in uploaded_files
            if len(item.getvalue()) > MAX_FILE_SIZE
        ]

        if oversized:
            st.warning(
                "Some files exceed the 10 MB limit: "
                + ", ".join(oversized)
            )

        st.caption(
            "To remove a file, use the remove (×) control on its upload entry. "
            "You can also clear the uploader and select another set."
        )

    else:
        st.info(
            "Your workspace is empty. Upload a PDF, DOCX, or TXT file to get started."
        )


# =========================================================
# ANALYZE ACTION
# =========================================================

st.markdown("")

can_analyze = (
    bool(uploaded_files)
    and online
    and model_available
    and any(
        len(item.getvalue()) <= MAX_FILE_SIZE
        for item in uploaded_files
    )
)

if st.button(
    "✦  Analyze documents with AI",
    type="primary",
    use_container_width=True,
    disabled=not can_analyze,
    help="Analyze the selected documents using your local Gemma 3 model.",
):
    pending = []

    for uploaded_file in uploaded_files:
        data = uploaded_file.getvalue()

        if len(data) > MAX_FILE_SIZE:
            continue

        digest = file_digest(data)

        if digest not in st.session_state.results:
            pending.append((uploaded_file.name, data, digest))

    if not pending:
        st.info("All valid selected files have already been analyzed.")
    else:
        progress = st.progress(0, text="Preparing your documents...")
        success_count = 0
        errors = []

        for index, (filename, data, digest) in enumerate(pending):
            progress.progress(
                index / len(pending),
                text=f"Analyzing {filename}...",
            )

            try:
                extracted_text = get_file_text(filename, data)

                result = analyze_document(
                    filename,
                    extracted_text,
                )

                st.session_state.results[digest] = result
                success_count += 1

            except requests.Timeout:
                errors.append(
                    f"{filename}: The AI took too long to respond. Try again."
                )

            except requests.RequestException:
                errors.append(
                    f"{filename}: Could not communicate with Ollama. "
                    "Check that Ollama is running."
                )

            except Exception as exc:
                errors.append(f"{filename}: {exc}")

        progress.progress(1.0, text="Analysis finished.")

        if success_count:
            st.success(
                f"Successfully analyzed {success_count} document(s)."
            )

        for error in errors:
            st.error(error)

        if success_count:
            st.rerun()


# =========================================================
# RESULTS
# =========================================================

st.markdown("")

st.markdown(
    """
    <div class="section-kicker">AI INSIGHTS</div>
    <div class="section-title">Analysis results</div>
    <div class="section-copy">
        Review the categories and summaries produced by your local AI model.
        Original uploads remain unchanged.
    </div>
    """,
    unsafe_allow_html=True,
)

if not st.session_state.results:
    with st.container(border=True):
        st.markdown(
            """
            <div style="text-align:center; padding:30px 10px;">
                <div style="font-size:40px; margin-bottom:12px;">✨</div>
                <div style="
                    font-family:Manrope,sans-serif;
                    font-size:18px;
                    font-weight:800;
                    color:#222944;
                    margin-bottom:8px;
                ">
                    Your AI insights will appear here
                </div>
                <div style="
                    color:#858da2;
                    font-size:13px;
                    line-height:1.8;
                ">
                    Upload your files and click Analyze documents with AI
                    to see real summaries and suggested categories.
                </div>
            </div>
            """,
            unsafe_allow_html=True,
        )

else:
    result_items = list(st.session_state.results.items())

    filter_options = ["All categories"] + sorted({
        result["category"]
        for _, result in result_items
    })

    filter_column, sort_column = st.columns([2, 1])

    with filter_column:
        selected_category = st.selectbox(
            "Filter by category",
            filter_options,
        )

    with sort_column:
        sort_order = st.selectbox(
            "Sort results",
            ["Original order", "Filename A–Z", "Category A–Z"],
        )

    visible_results = result_items

    if selected_category != "All categories":
        visible_results = [
            (digest, result)
            for digest, result in visible_results
            if result["category"] == selected_category
        ]

    if sort_order == "Filename A–Z":
        visible_results.sort(
            key=lambda item: item[1]["suggested_filename"].lower()
        )
    elif sort_order == "Category A–Z":
        visible_results.sort(
            key=lambda item: item[1]["category"].lower()
        )

    st.caption(
        f"Showing {len(visible_results)} of {len(result_items)} analyzed documents"
    )

    for index, (digest, result) in enumerate(visible_results, start=1):
        with st.container(border=True):
            st.markdown(
                f"""
                <div class="category-pill">
                    {html.escape(result["category"])}
                </div>
                <div style="
                    font-family:Manrope,sans-serif;
                    font-size:17px;
                    font-weight:800;
                    color:#202742;
                    margin-bottom:5px;
                    overflow-wrap:anywhere;
                ">
                    {html.escape(result["original_name"])}
                </div>
                <div class="result-summary">
                    {html.escape(result["summary"])}
                </div>
                <div style="
                    font-size:11px;
                    color:#828ba1;
                    font-weight:700;
                    margin-top:16px;
                    margin-bottom:7px;
                ">
                    SUGGESTED ORGANIZED FILENAME
                </div>
                <div class="result-filename">
                    {html.escape(result["suggested_filename"])}
                </div>
                """,
                unsafe_allow_html=True,
            )

            edit_key = f"edit_{digest}"

            with st.expander("Review or edit this result"):
                edited_category = st.selectbox(
                    "Category",
                    CATEGORIES,
                    index=CATEGORIES.index(result["category"]),
                    key=f"category_{edit_key}",
                )

                edited_summary = st.text_area(
                    "Summary",
                    value=result["summary"],
                    key=f"summary_{edit_key}",
                    height=110,
                )

                edited_filename = st.text_input(
                    "Suggested filename",
                    value=result["suggested_filename"],
                    key=f"filename_{edit_key}",
                )

                if st.button(
                    "Save changes",
                    key=f"save_{edit_key}",
                    use_container_width=True,
                ):
                    filename = safe_filename(edited_filename)

                    if not Path(filename).suffix:
                        original_extension = Path(
                            result["original_name"]
                        ).suffix.lower()
                        filename += original_extension

                    result["category"] = edited_category
                    result["summary"] = edited_summary.strip()
                    result["suggested_filename"] = filename

                    st.session_state.results[digest] = result

                    st.success("Your changes have been saved.")
                    st.rerun()


# =========================================================
# EXPORT
# =========================================================

st.markdown("")

st.markdown(
    """
    <div class="section-kicker">TAKE YOUR FILES WITH YOU</div>
    <div class="section-title">Export organized documents</div>
    <div class="section-copy">
        Download category-wise document copies together with a JSON report
        containing the AI-generated summaries and filenames.
    </div>
    """,
    unsafe_allow_html=True,
)

if st.session_state.results:
    with st.container(border=True):
        export_col, details_col = st.columns([1.2, 1])

        with export_col:
            st.markdown("### 📦 Your organization archive")
            st.markdown(
                "A ZIP archive containing organized document copies "
                "and `organization_report.json`."
            )

            archive = build_zip(
                st.session_state.results,
                st.session_state.file_data,
            )

            st.download_button(
                "⬇️ Download organized ZIP",
                data=archive,
                file_name="lazydesk_organized_documents.zip",
                mime="application/zip",
                type="primary",
                use_container_width=True,
            )

        with details_col:
            st.markdown("### Archive contents")

            st.markdown(
                f"- **{len(st.session_state.results)}** analyzed documents\n"
                f"- **{category_count}** categories in your workspace\n"
                "- Organized copies in category folders\n"
                "- JSON report with summaries and filenames"
            )

    st.caption(
        "The archive contains copies of your uploaded files. "
        "Your original documents are not modified."
    )

else:
    st.info(
        "Analyze at least one document to enable the organized ZIP download."
    )


# =========================================================
# FOOTER
# =========================================================

st.markdown(
    """
    <div class="footer">
        <strong>✦ LazyDesk AI</strong><br>
        Local-first document organization powered by Gemma 3 + Ollama.<br>
        Built to make everyday document work simpler.
    </div>
    """,
    unsafe_allow_html=True,
)